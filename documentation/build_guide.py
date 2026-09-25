from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "documentation" / "ECDAT_Beginner_Guide.docx"

INK = "102023"
GREEN = "247A61"
PALE = "EAF3EE"
LIGHT = "F4F6F2"
MID = "647370"
LINE = "D9D9D9"
WHITE = "FFFFFF"


def set_cell_fill(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_border(cell, color=LINE):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=120, start=140, bottom=120, end=140):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_table(table, header=True):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_border(cell)
            set_cell_margins(cell)
            if header and r_index == 0:
                set_cell_fill(cell, INK)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.color.rgb = RGBColor(255, 255, 255)
                        run.font.bold = True
            elif r_index % 2 == 0:
                set_cell_fill(cell, LIGHT)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
    for row_data in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row_data):
            cells[i].text = str(value)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    shade_table(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_bullet(doc, text, level=0):
    paragraph = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    paragraph.add_run(text)
    return paragraph


def reset_number(doc):
    doc._ecdat_number = 0


def add_number(doc, title, explanation):
    doc._ecdat_number = getattr(doc, "_ecdat_number", 0) + 1
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.18)
    paragraph.paragraph_format.first_line_indent = Inches(-0.18)
    run = paragraph.add_run(f"{doc._ecdat_number}.  {title} ")
    run.bold = True
    paragraph.add_run(explanation)
    return paragraph


def page_break(doc):
    doc.add_page_break()


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.14

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title.font.size = Pt(32)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(14)
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for name, size, before, after in (("Heading 1", 21, 18, 8), ("Heading 2", 14, 13, 5), ("Heading 3", 11, 9, 3)):
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for list_name in ("List Bullet", "List Bullet 2", "List Number"):
        styles[list_name].font.name = "Aptos"
        styles[list_name].font.size = Pt(10.2)
        styles[list_name].paragraph_format.space_after = Pt(4)

    footer = section.footer.paragraphs[0]
    footer.text = "ECDAT beginner guide   |   Smart India Hackathon 2026"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in footer.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(MID)

    # Cover
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(52)
    eyebrow = doc.add_paragraph()
    eyebrow.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = eyebrow.add_run("PROJECT GUIDE  |  SIH26164")
    run.font.name = "Aptos"
    run.font.size = Pt(10)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(GREEN)

    cover_title = doc.add_paragraph("Enterprise Cryptographic Discovery and Assessment Tool", style="Title")
    cover_ppr = cover_title._p.get_or_add_pPr()
    cover_border = cover_ppr.find(qn("w:pBdr"))
    if cover_border is not None:
        cover_ppr.remove(cover_border)
    subtitle = doc.add_paragraph("A beginner friendly guide to ECDAT and its VS Code extension")
    subtitle.paragraph_format.space_after = Pt(28)
    for run in subtitle.runs:
        run.font.name = "Aptos Display"
        run.font.size = Pt(17)
        run.font.color.rgb = RGBColor.from_string(MID)

    intro = doc.add_paragraph()
    intro.add_run("What this document does. ").bold = True
    intro.add_run(
        "It explains the problem, the complete scanning pipeline, the privacy model, and live local scanning inside VS Code. It assumes no background in cryptography or security tools."
    )
    intro.paragraph_format.space_after = Pt(20)

    add_table(doc, ["Project", "Value"], [
        ("Name", "ECDAT"),
        ("Problem statement", "SIH26164"),
        ("Organization", "National Technical Research Organisation"),
        ("Main output", "Findings dashboard and Cryptography Bill of Materials"),
        ("Current status", "Working multi-surface scanner and VS Code extension"),
    ], [1.75, 4.8])

    doc.add_paragraph("The main idea in one sentence", style="Heading 2")
    statement = doc.add_paragraph()
    statement.add_run(
        "ECDAT scans repositories, binaries, container images and cryptographic artifacts, explains what is risky, and recommends safer or post-quantum alternatives."
    ).bold = True

    page_break(doc)

    # Basics
    doc.add_paragraph("1  The problem ECDAT solves", style="Heading 1")
    doc.add_paragraph(
        "Modern applications use encryption, digital signatures, hashing, certificates and cryptographic libraries in many places. Over time, teams may lose track of where these components are used. Some choices become outdated, while public key algorithms such as RSA and elliptic curve cryptography also need long term quantum migration planning."
    )
    doc.add_paragraph("A simple example", style="Heading 2")
    doc.add_paragraph(
        "Imagine a payment application with one secure AES encryption function, an old SHA-1 checksum, an RSA certificate and a private key accidentally committed to the repository. A developer can search manually, but the result may be incomplete and difficult to explain. ECDAT performs the inventory consistently and reports each finding with its location, risk and recommended action."
    )

    doc.add_paragraph("Important terms", style="Heading 2")
    add_table(doc, ["Term", "Plain English meaning"], [
        ("Cryptography", "Methods used to protect data, prove identity or detect changes."),
        ("Algorithm", "A method such as AES, RSA, SHA-256 or MD5."),
        ("Key", "A secret or public value used by an encryption or signing algorithm."),
        ("Certificate", "A digital identity document that connects a public key to a system or organization."),
        ("HSM", "A protected device used to generate, store and use cryptographic keys."),
        ("CBOM", "A structured inventory of the cryptography used by a software project."),
        ("Post quantum cryptography", "Algorithms designed to resist attacks from future large quantum computers."),
    ], [1.65, 4.9])

    doc.add_paragraph("What ECDAT is and is not", style="Heading 2")
    add_table(doc, ["ECDAT does", "ECDAT does not"], [
        ("Scan source, binaries, containers, certificates and keys", "Execute uploaded application code or binaries"),
        ("Find likely cryptographic usage", "Claim that every match is a confirmed vulnerability"),
        ("Apply transparent prototype rules", "Use a hidden AI model to invent risk scores"),
        ("Redact detected private key material", "Display or export secret key contents"),
        ("Support migration planning", "Predict the exact arrival date of a cryptographically relevant quantum computer"),
    ], [3.25, 3.25])

    # Current product
    doc.add_paragraph("2  How the complete scanner works", style="Heading 1")
    doc.add_paragraph("The full workflow is deliberately easy to demonstrate.")
    reset_number(doc)
    add_number(doc, "Choose an input.", "Upload a repository ZIP or TAR, container-image archive, binary, certificate or key, or run the included demonstration project.")
    add_number(doc, "Set the context.", "Choose data sensitivity, migration complexity and an assumed quantum threat timeline.")
    add_number(doc, "Start the scan.", "ECDAT reads supported files without running the uploaded code.")
    add_number(doc, "Review findings.", "The dashboard shows the detected item, file, line, risk, quantum risk and recommendation.")
    add_number(doc, "Export the inventory.", "Download the CBOM as JSON for later review or integration.")

    doc.add_paragraph("What the scanner detects", style="Heading 2")
    add_table(doc, ["Category", "Examples", "Why it matters"], [
        ("Algorithms", "AES, RSA, ECDSA, SHA-1, SHA-256, MD5, ChaCha20", "Shows modern, legacy and quantum vulnerable choices."),
        ("Libraries", "OpenSSL, PyCryptodome, Python cryptography, Node Crypto, Web Crypto, libsodium and Bouncy Castle", "Identifies dependencies that must be maintained."),
        ("Keys", "PEM, DER, public, private and PKCS number 12 containers", "OpenSSL validates structures without revealing key contents."),
        ("Certificates", "PEM, CRT, CER and DER files", "OpenSSL extracts subject, expiry, signature and public-key strength."),
        ("HSM usage", "PKCS number 11 and HSM references", "Shows where protected hardware key storage may be used."),
        ("Containers", "Syft package inventory and Trivy vulnerability results", "Finds cryptographic packages inside image archives."),
    ], [1.15, 2.5, 2.85])

    doc.add_paragraph("How risk is calculated", style="Heading 2")
    doc.add_paragraph(
        "Classical risk uses readable rules. For example, MD5 and SHA-1 are high risk for security-sensitive use, a committed private key is critical, and modern authenticated encryption such as AES-256-GCM is low risk under the prototype rules. These ratings are guidance for review, not an official government or standards-body score."
    )

    doc.add_paragraph("Quantum timeline model", style="Heading 2")
    doc.add_paragraph(
        "ECDAT uses the Mosca-style inequality X + Y > Z. X is the number of years the data must remain secure. Y is the time needed to migrate the system. Z is an assumed quantum threat timeline. When data lifetime plus migration time is greater than the threat timeline, planning is urgent. This logic is applied to quantum-vulnerable public key cryptography, not blindly to every symmetric algorithm or hash."
    )

    # Privacy
    doc.add_paragraph("3  Privacy and data handling", style="Heading 1")
    doc.add_paragraph(
        "The scanner is designed to avoid retaining uploaded inputs. Archives and artifacts are written to a temporary directory, validated, scanned and then removed. Source code and unknown binaries are inspected but never executed."
    )
    add_table(doc, ["Data", "Current behavior", "Recommended production behavior"], [
        ("Uploaded archives and artifacts", "Temporary local processing, then deletion", "Keep processing local or on an isolated customer-controlled worker"),
        ("Private key material", "Never displayed; evidence is redacted", "Block storage completely and alert the user to rotate the key"),
        ("Scan summary", "Stored in local SQLite", "Make retention configurable and add automatic deletion"),
        ("File paths and safe evidence", "Stored with scan findings", "Allow evidence-free mode for sensitive organizations"),
        ("External services", "No source code is sent to an external AI service", "Preserve this local-first rule"),
    ], [1.4, 2.45, 2.65])

    doc.add_paragraph("Why a VS Code extension improves privacy", style="Heading 2")
    doc.add_paragraph(
        "A VS Code extension can run the same detection rules directly inside the developer's workspace. The source stays on the developer's machine, feedback appears while code is being written, and only an optional sanitized report needs to be exported. This removes the main concern behind asking users to upload their repository to a website."
    )

    doc.add_paragraph("Security controls already present", style="Heading 2")
    for item in (
        "ZIP and TAR path traversal protection prevents files from escaping the temporary extraction folder.",
        "Upload, extracted-size and file-count limits reduce resource abuse.",
        "Large or unsupported files are skipped instead of being read blindly.",
        "Generated dependency folders such as node_modules and virtual environments are skipped.",
        "OpenSSL parsing validates certificates and key structures while redacting private material.",
        "Syft and Trivy inspect container packages without running the image.",
        "Normal users receive useful errors instead of raw Python stack traces.",
    ):
        add_bullet(doc, item)

    # VS Code plan
    doc.add_paragraph("4  VS Code live scanning extension", style="Heading 1")
    doc.add_paragraph(
        "An initial extension is now included in the repository. It turns ECDAT from a scan-on-demand web tool into continuous developer feedback while preserving the existing scanner, risk and recommendation logic."
    )

    doc.add_paragraph("Developer experience", style="Heading 2")
    reset_number(doc)
    add_number(doc, "Open a project in VS Code.", "The extension recognizes supported source and configuration files.")
    add_number(doc, "Edit or save a file.", "Only the changed file is scanned for fast feedback.")
    add_number(doc, "See an inline warning.", "Risky cryptography is underlined and appears in the Problems panel.")
    add_number(doc, "Open the explanation.", "The developer sees the reason, quantum context and recommended action.")
    add_number(doc, "Run a workspace scan.", "A command scans the complete project and opens the ECDAT dashboard view.")
    add_number(doc, "Export the CBOM.", "The extension writes a sanitized JSON report selected by the user.")

    doc.add_paragraph("Simple architecture", style="Heading 2")
    add_table(doc, ["Part", "Responsibility", "Suggested technology"], [
        ("VS Code extension", "Commands, settings, diagnostics and Problems panel integration", "TypeScript and VS Code Extension API"),
        ("Local scanner", "File discovery and cryptographic pattern detection", "Reuse the current Python engine first"),
        ("Local bridge", "Pass file text and results between the extension and scanner", "A local process using JSON over standard input and output"),
        ("Webview dashboard", "Workspace summary, filters, details and CBOM export", "Reuse the React dashboard design"),
        ("Rules configuration", "Control severity, exclusions and organization policy", "Local JSON settings"),
    ], [1.3, 3.2, 2.0])

    doc.add_paragraph("Live scanning rules", style="Heading 2")
    add_bullet(doc, "Scan the active file on save, not on every keystroke.")
    add_bullet(doc, "Debounce repeated saves and cancel stale scans.")
    add_bullet(doc, "Scan only supported files and respect workspace exclusions.")
    add_bullet(doc, "Never send source text over the network.")
    add_bullet(doc, "Show a short diagnostic message; keep the full explanation in a detail panel.")
    add_bullet(doc, "Provide a clear command to rescan the whole workspace.")

    # Implementation
    doc.add_paragraph("5  Extension implementation status", style="Heading 1")
    doc.add_paragraph(
        "The first working version covers the full local scan workflow. The team can use this table to understand what is already present and what must be validated before packaging."
    )
    add_table(doc, ["Stage", "Deliverable", "Status"], [
        ("1", "ECDAT commands and extension activation", "Implemented"),
        ("2", "Current-file scanning through the local Python bridge", "Implemented"),
        ("3", "Editor and Problems-panel diagnostics", "Implemented"),
        ("4", "Save-triggered scanning", "Implemented"),
        ("5", "Workspace scan and summary webview", "Implemented"),
        ("6", "CBOM export and user settings", "Implemented"),
        ("7", "Installation package and user testing", "Validate with the team"),
    ], [0.65, 2.7, 3.15])

    doc.add_paragraph("Suggested division of work", style="Heading 2")
    add_table(doc, ["Area", "Main tasks"], [
        ("Scanner and security", "Maintain detection rules, risk logic, redaction and archive safety."),
        ("VS Code extension", "Implement commands, diagnostics, local bridge and extension settings."),
        ("Dashboard and demo", "Adapt the React summary view, prepare demo cases and check usability."),
        ("Testing and documentation", "Test scanner behavior, extension workflow, privacy controls and setup instructions."),
    ], [1.75, 4.75])

    doc.add_paragraph("Minimum acceptance checks", style="Heading 2")
    for item in (
        "A saved file with SHA-1 receives a warning on the correct line.",
        "AES-256-GCM is inventoried without an urgent classical-risk warning.",
        "RSA receives a quantum migration warning based on selected assumptions.",
        "Private-key evidence is always replaced with a redacted message.",
        "Excluded folders are not scanned.",
        "No source code leaves the machine during normal use.",
        "The exported CBOM contains real findings rather than demonstration numbers.",
    ):
        add_bullet(doc, item)

    page_break(doc)

    # Demo and talking points
    doc.add_paragraph("6  How to explain and demonstrate ECDAT", style="Heading 1")
    doc.add_paragraph("A short explanation for a new teammate", style="Heading 2")
    doc.add_paragraph(
        "ECDAT is like an inventory and warning system for cryptography in software. It searches the project for encryption, hashes, certificates, keys and crypto libraries. It then explains which items need attention now and which public key items need quantum migration planning. It does not run the project and it hides secret key material."
    )

    doc.add_paragraph("Five minute web demo", style="Heading 2")
    reset_number(doc)
    add_number(doc, "Open the dashboard.", "Point out that scanning is local and rule based.")
    add_number(doc, "Run the sample project.", "The current sample scans seven files and produces fifteen findings.")
    add_number(doc, "Open the private-key finding.", "Show that evidence is redacted and the action recommends removal and rotation.")
    add_number(doc, "Open an RSA finding.", "Explain the selected X, Y and Z assumptions and the migration margin.")
    add_number(doc, "Export the CBOM.", "Show that the report is generated from the real scan result.")

    doc.add_paragraph("Final takeaway", style="Heading 2")
    doc.add_paragraph(
        "The ECDAT repository provides the complete discovery-to-report workflow across repositories, binaries, container images, certificates and keys, plus a local VS Code extension for continuous feedback."
    )

    page_break(doc)
    doc.add_paragraph("Questions the audience may ask", style="Heading 2")
    add_table(doc, ["Question", "Short answer"], [
        ("Does the server keep my input?", "The scanner deletes temporary archives, binaries and extracted files after scanning. It stores the sanitized result locally in SQLite."),
        ("Is this an AI model?", "No. Discovery, risk and recommendations use explicit rules that developers can read and test."),
        ("Can it guarantee that an application is secure?", "No. It creates an inventory and prioritizes review; security specialists must validate the context."),
        ("Why add a VS Code extension?", "It gives immediate local feedback and avoids uploading source code to a web service."),
        ("Does quantum risk apply to AES too?", "Not in the same way. The Mosca timeline is applied to public key algorithms affected by Shor's algorithm."),
        ("Which engines are used?", "Source rules plus an OpenSSL backend, binary fingerprints, Syft package inventory and Trivy container-vulnerability analysis."),
    ], [2.25, 4.25])

    # Apply typography consistently.
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            if not run.font.name:
                run.font.name = "Aptos"
            run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), run.font.name)
            run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), run.font.name)

    core = doc.core_properties
    core.title = "ECDAT Beginner Guide"
    core.subject = "ECDAT web prototype and VS Code extension plan"
    core.author = "ECDAT Team"
    core.keywords = "ECDAT, cryptography, CBOM, VS Code, SIH26164"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
